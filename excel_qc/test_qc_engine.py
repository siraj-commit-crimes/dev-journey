"""Tests for qc_engine using a fake Claude client (no API key or internet needed).

Run with:  python -m pytest excel_qc
"""

import json
from io import BytesIO
from types import SimpleNamespace

import pandas as pd
from openpyxl import load_workbook

import qc_engine


class FakeMessages:
    """Pretends to be client.beta.messages and answers like Claude would."""

    def __init__(self):
        self.requests = []

    def create(self, **request):
        self.requests.append(request)
        payload = json.loads(request["messages"][0]["content"])
        if "topics" in payload:
            unified = {"Customer Service": "Customer Service", "customer service": "Customer Service",
                       "خدمة العملاء": "Customer Service", "Prices": "Pricing"}
            answer = {"mapping": [{"original": t["label"], "unified": unified.get(t["label"], t["label"])}
                                  for t in payload["topics"]]}
        else:
            answer = {"results": [self.judge(row, payload["brand_list"]) for row in payload["rows"]]}
        block = SimpleNamespace(type="text", text=json.dumps(answer))
        return SimpleNamespace(stop_reason="end_turn", content=[block])

    @staticmethod
    def judge(row, brand_list):
        sarcastic = "👏" in row["text"]
        mentioned = [b["name"] for b in brand_list
                     if any(n.lower() in row["text"].lower() for n in [b["name"], *b["aliases"]])]
        target = row["brand"] or (mentioned[0] if mentioned else "")
        return {
            "id": row["id"],
            "gender": "Female" if row["author"] in ("Noura", "سارة") else "Male",
            "gender_reason": "name",
            "sentiment": "Negative" if sarcastic else "Positive",
            "sarcasm": sarcastic,
            "sentiment_reason": "test",
            "target_brand": target,
            "brands_mentioned": mentioned,
            "brand_sentiments": [{"brand": b, "sentiment": "Positive" if i == 0 else "Negative"}
                                 for i, b in enumerate(mentioned)],
            "relevant": bool(mentioned) or not (brand_list or row["brand"]),
            "topic_fits": row["topic"] != "Pricing",
            "suggested_topic": "Customer Service",
        }


def fake_client():
    return SimpleNamespace(beta=SimpleNamespace(messages=FakeMessages()))


def sample_df():
    return pd.DataFrame({
        "Author Name": ["Noura", "Khalid", "سارة", "Fahad"],
        "Screen Name": ["@noura_99", "@abu_khalid", "@sara", "@fahad"],
        "Content": ["خدمة ممتازة 😍", "Great service, waited 3 hours 👏", "شكرا", "Love it"],
        "Gender": ["F", "male", "ذكر", "Male"],
        "Sentiment": ["Positive", "Positive", "إيجابي", "positive"],
        "Topic": ["Customer Service", "customer service", "خدمة العملاء", "Prices"],
    })


def test_guess_columns():
    guesses = qc_engine.guess_columns(sample_df().columns)
    assert guesses == {
        "text": "Content", "handle": "Screen Name", "author": "Author Name",
        "gender": "Gender", "sentiment": "Sentiment", "brand": None, "topic": "Topic",
    }


def test_normalize_labels():
    assert qc_engine.normalize_gender("F") == "Female"
    assert qc_engine.normalize_gender("ذكر") == "Male"
    assert qc_engine.normalize_gender(None) == "Unknown"
    assert qc_engine.normalize_sentiment("سلبي") == "Negative"


def test_run_qc_flags_mismatches():
    df = sample_df()
    columns = qc_engine.guess_columns(df.columns)
    result, summary, topic_map, brand_df = qc_engine.run_qc(df, columns, fake_client(), batch_size=2)

    assert list(result["QC_Status"]) == ["OK", "Check", "Check", "Check"]
    assert result.loc[1, "QC_Flags"] == "Sentiment"       # sarcasm caught
    assert result.loc[1, "QC_Sarcasm"] == "Yes"
    assert result.loc[2, "QC_Flags"] == "Gender"          # سارة marked as male
    assert result.loc[3, "QC_Flags"] == "Topic"
    assert set(result["Unified_Topic"]) == {"Customer Service", "Pricing"}
    assert len(topic_map) == 4

    counts = dict(zip(summary["Check"], summary["Count"]))
    assert counts["Topics before unifying"] == 4
    assert counts["Topics after unifying"] == 2
    # No brands given: no brand columns and no Irrelevant flags.
    assert "QC_Relevant" not in result
    assert brand_df.empty


def test_failed_batch_marks_rows_as_error():
    client = fake_client()

    def broken(**request):
        raise RuntimeError("boom")

    client.beta.messages.create = broken
    df = sample_df()
    columns = qc_engine.guess_columns(df.columns)
    columns["topic"] = None  # skip topic unifying so only the row check fails
    result, summary, _, _ = qc_engine.run_qc(df, columns, client)
    assert set(result["QC_Status"]) == {"Error"}
    assert "boom" in result.loc[0, "QC_Flags"]


def test_excel_output_is_highlighted():
    df = sample_df()
    columns = qc_engine.guess_columns(df.columns)
    result, summary, topic_map, _ = qc_engine.run_qc(df, columns, fake_client())
    data = qc_engine.to_excel(result, summary, topic_map, columns)

    book = load_workbook(BytesIO(data))
    assert book.sheetnames == ["Summary", "QC Results", "Topic Mapping"]
    sheet = book["QC Results"]
    headers = [c.value for c in sheet[1]]
    sentiment_cell = sheet.cell(3, headers.index("Sentiment") + 1)  # row 2 of data (sarcasm)
    assert sentiment_cell.fill.start_color.rgb.endswith("F8CBAD")


def test_haiku_has_no_fallback():
    client = fake_client()
    df = sample_df().head(1)
    columns = qc_engine.guess_columns(df.columns)
    qc_engine.run_qc(df, columns, client, model="claude-haiku-5-5")
    assert all("fallbacks" not in r for r in client.beta.messages.requests)
    qc_engine.run_qc(df, columns, client, model="claude-opus-5-5")
    assert client.beta.messages.requests[-1]["fallbacks"] == "default"


def test_parse_brands():
    brands = qc_engine.parse_brands("Client: STC, اس تي سي, stc_ksa\n\ncompetitor: Mobily\nZain, زين")
    assert brands == [
        {"name": "STC", "role": "Client", "aliases": ["اس تي سي", "stc_ksa"]},
        {"name": "Mobily", "role": "Competitor", "aliases": []},
        {"name": "Zain", "role": "Competitor", "aliases": ["زين"]},
    ]


def test_brands_and_competitors():
    df = pd.DataFrame({
        "Content": ["STC is way better than Mobily 👌", "الطقس حار اليوم", "Mobily down again"],
        "Sentiment": ["Positive", "Neutral", "Negative"],
    })
    brands = qc_engine.parse_brands("Client: STC, اس تي سي\nCompetitor: Mobily, موبايلي")
    columns = qc_engine.guess_columns(df.columns)
    client = fake_client()
    result, summary, _, brand_df = qc_engine.run_qc(df, columns, client, brands=brands)

    assert client.beta.messages.requests[0]["messages"][0]["content"].count("Mobily") >= 1
    assert list(result["QC_Target_Brand"]) == ["STC", "", "Mobily"]
    assert list(result["QC_Relevant"]) == ["Yes", "No", "Yes"]
    assert result.loc[0, "QC_Sentiment_By_Brand"] == "STC: Positive | Mobily: Negative"
    assert "Irrelevant" in result.loc[1, "QC_Flags"]

    by_brand = brand_df.set_index("Brand")
    assert by_brand.loc["Mobily", "Mentions"] == 2
    assert by_brand.loc["STC", "Net_Sentiment"] == "100%"

    counts = dict(zip(summary["Check"], summary["Count"]))
    assert counts["Irrelevant posts (not about the brands)"] == 1

    data = qc_engine.to_excel(result, summary, pd.DataFrame(), columns, brand_df)
    assert "Brand Sentiment" in load_workbook(BytesIO(data)).sheetnames


def test_brand_column_sets_target():
    df = pd.DataFrame({
        "Brand": ["STC", "Mobily"],
        "Content": ["STC is better than Mobily", "STC is better than Mobily"],
    })
    columns = qc_engine.guess_columns(df.columns)
    assert columns["brand"] == "Brand"
    client = fake_client()
    result, *_ = qc_engine.run_qc(df, columns, client)
    sent_rows = json.loads(client.beta.messages.requests[0]["messages"][0]["content"])["rows"]
    assert [r["brand"] for r in sent_rows] == ["STC", "Mobily"]
    assert list(result["QC_Target_Brand"]) == ["STC", "Mobily"]
