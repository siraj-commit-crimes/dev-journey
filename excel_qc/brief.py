"""Read the client's brief (Word, PDF or text) into plain text for Claude."""

from io import BytesIO

import docx
from pypdf import PdfReader

MAX_CHARS = 60_000  # about 15k tokens; a brief is usually a few pages


def read_brief(name, data):
    """Return the brief's text. Raises ValueError for files it can't read."""
    lower = name.lower()
    if lower.endswith(".docx"):
        text = _read_docx(data)
    elif lower.endswith(".pdf"):
        reader = PdfReader(BytesIO(data))
        text = "\n".join(page.extract_text() or "" for page in reader.pages)
    elif lower.endswith((".txt", ".md")):
        text = data.decode("utf-8-sig", errors="replace")
    else:
        raise ValueError("Upload a Word (.docx), PDF or text file.")
    text = clean_text(text)
    if not text:
        raise ValueError("No text found in the brief. If it's a scanned PDF, paste the text instead.")
    return text


def _read_docx(data):
    document = docx.Document(BytesIO(data))
    lines = [p.text for p in document.paragraphs]
    # Briefs often keep brand and topic lists in tables.
    for table in document.tables:
        for row in table.rows:
            cells = [c.text.strip() for c in row.cells]
            lines.append(" | ".join(dict.fromkeys(c for c in cells if c)))
    return "\n".join(lines)


def clean_text(text):
    lines = [line.strip() for line in str(text or "").splitlines()]
    text = "\n".join(line for line in lines if line)
    return text[:MAX_CHARS]
