"""Tests for the review step: review tables, final values and the Final Data sheet."""

from io import BytesIO

import pandas as pd
from openpyxl import load_workbook

import offline_engine as offline
import qc_engine


def run_sample():
    df = pd.DataFrame({
        "Author": ["Noura", "Khalid", "Ali"],
        "Content": ["STC ممتازة 😍", "Great service STC, waited 3 hours 👏", "الطقس حار اليوم"],
        "Gender": ["Male", "Male", "Male"],
        "Sentiment": ["Positive", "Positive", "Neutral"],
        "Topic": ["Service", "service", "Other"],
    })
    columns = qc_engine.guess_columns(df.columns)
    brands = qc_engine.parse_brands("Client: STC")
    result, summary, topic_map, brand_df = qc_engine.run_qc(
        df, columns, offline.OfflineClient(), model="offline", brands=brands
    )
    return df, columns, result, summary, topic_map, brand_df


def test_review_tables():
    _, columns, result, *_ = run_sample()
    gender = qc_engine.review_table(result, columns, "Gender")
    assert list(gender.index) == [0]
    assert gender.loc[0, ["Row", "Original", "Suggested"]].tolist() == [2, "Male", "Female"]

    sentiment = qc_engine.review_table(result, columns, "Sentiment")
    assert list(sentiment.index) == [1]
    assert sentiment.loc[1, "Reason"].startswith("Sarcasm.")

    irrelevant = qc_engine.review_table(result, columns, "Irrelevant")
    assert list(irrelevant.index) == [2]

    assert qc_engine.review_options(gender, "Gender")[:3] == ["Male", "Female", "Unknown"]


def test_build_final_uses_decisions_and_keeps_originals():
    df, columns, result, *_ = run_sample()
    decisions = {"Gender": {0: "Female"}, "Irrelevant": {2: "Irrelevant"}}
    out, final = qc_engine.build_final(result, columns, decisions)

    assert list(final.columns) == list(df.columns) + ["Relevant"]
    assert list(final["Gender"]) == ["Female", "Male", "Male"]
    assert list(final["Sentiment"]) == ["Positive", "Positive", "Neutral"]  # not reviewed: kept
    assert list(final["Relevant"]) == ["Relevant", "Relevant", "Irrelevant"]
    assert list(final["Topic"]) == ["Service", "Service", "Other"]  # unified topics
    assert list(out["Review_Status"]) == ["Reviewed", "Not reviewed", "Reviewed"]

    counts = qc_engine.review_counts(result, columns, decisions)
    assert counts == {"Gender": (1, 1), "Sentiment": (0, 1), "Irrelevant": (1, 1)}


def test_topic_map_edit_changes_final_topic():
    _, columns, result, *_ = run_sample()
    edited = qc_engine.apply_topic_map(result, columns, {"Service": "Customer Service", "service": "Customer Service"})
    _, final = qc_engine.build_final(edited, columns, {})
    assert list(final["Topic"]) == ["Customer Service", "Customer Service", "Other"]


def test_final_data_sheet_marks_changes():
    _, columns, result, summary, topic_map, brand_df = run_sample()
    out, final = qc_engine.build_final(result, columns, {"Gender": {0: "Female"}})
    data = qc_engine.to_excel(out, summary, topic_map, columns, brand_df, final)
    book = load_workbook(BytesIO(data))
    assert book.sheetnames[:3] == ["Summary", "Final Data", "QC Results"]
    sheet = book["Final Data"]
    headers = [c.value for c in sheet[1]]
    gender_col = headers.index("Gender") + 1
    assert sheet.cell(2, gender_col).value == "Female"
    assert sheet.cell(2, gender_col).fill.start_color.rgb.endswith("E2EFDA")   # changed
    assert not sheet.cell(3, gender_col).fill.start_color.rgb.endswith("E2EFDA")  # unchanged


def test_clean_handles_missing_values():
    assert qc_engine.clean(pd.NA) == ""
    assert qc_engine.clean(float("nan")) == ""
    assert qc_engine.clean(None) == ""
    assert qc_engine.clean(" x ") == "x"
