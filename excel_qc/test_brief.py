"""Tests for the client brief: reading files, sending it to Claude, filling in Setup."""

import json
from io import BytesIO
from types import SimpleNamespace

import docx
import pandas as pd
import pytest

import brief
import qc_engine


def make_docx():
    document = docx.Document()
    document.add_heading("STC social listening brief", 1)
    document.add_paragraph("Relevant: posts about STC mobile, fiber and the stc pay app.")
    document.add_paragraph("Exclude job posts and football sponsorship news.")
    table = document.add_table(rows=2, cols=2)
    table.cell(0, 0).text, table.cell(0, 1).text = "Competitor", "Mobily"
    table.cell(1, 0).text, table.cell(1, 1).text = "Competitor", "زين"
    buffer = BytesIO()
    document.save(buffer)
    return buffer.getvalue()


def test_read_docx_includes_tables():
    text = brief.read_brief("Brief.DOCX", make_docx())
    assert "Exclude job posts" in text
    assert "Competitor | Mobily" in text
    assert "Competitor | زين" in text


def test_read_text_and_errors():
    assert brief.read_brief("b.txt", "  Relevant: STC\n\n\nNews is Neutral ".encode()) == "Relevant: STC\nNews is Neutral"
    with pytest.raises(ValueError, match="Word"):
        brief.read_brief("brief.xlsx", b"x")
    with pytest.raises(ValueError, match="No text"):
        brief.read_brief("empty.txt", b"   \n ")


class RecordingMessages:
    """Fake Claude that records requests and answers every row as irrelevant."""

    def __init__(self):
        self.requests = []

    def create(self, **request):
        self.requests.append(request)
        content = request["messages"][0]["content"]
        if request["system"] == qc_engine.SETUP_SYSTEM_PROMPT:
            answer = {
                "brands": [
                    {"name": "STC", "role": "Client", "aliases": ["اس تي سي", " stc_ksa ", ""]},
                    {"name": "Mobily", "role": "Competitor", "aliases": []},
                    {"name": " ", "role": "Competitor", "aliases": []},
                ],
                "topics": ["Network", "Pricing", "Network"],
            }
        else:
            rows = json.loads(content)["rows"]
            answer = {"results": [{
                "id": r["id"], "gender": "Male", "gender_reason": "", "sentiment": "Neutral",
                "sarcasm": False, "sentiment_reason": "Brief: news is Neutral", "target_brand": "",
                "brands_mentioned": [], "brand_sentiments": [], "relevant": False,
                "topic_fits": True, "suggested_topic": "",
            } for r in rows]}
        block = SimpleNamespace(type="text", text=json.dumps(answer, ensure_ascii=False))
        return SimpleNamespace(stop_reason="end_turn", content=[block])


def fake_client():
    return SimpleNamespace(beta=SimpleNamespace(messages=RecordingMessages()))


def test_brief_is_sent_and_turns_on_relevance():
    df = pd.DataFrame({"Content": ["STC is hiring engineers"], "Sentiment": ["Neutral"]})
    columns = qc_engine.guess_columns(df.columns)
    client = fake_client()
    result, *_ = qc_engine.run_qc(df, columns, client, brief="Exclude job posts.")

    payload = json.loads(client.beta.messages.requests[0]["messages"][0]["content"])
    assert payload["client_brief"] == "Exclude job posts."
    assert "Client brief" in client.beta.messages.requests[0]["system"]
    # A brief alone (no brand list) is enough to check relevance.
    assert result.loc[0, "QC_Relevant"] == "No"
    assert result.loc[0, "QC_Flags"] == "Irrelevant"


def test_no_brief_no_relevance_column():
    df = pd.DataFrame({"Content": ["hello"]})
    columns = qc_engine.guess_columns(df.columns)
    client = fake_client()
    result, *_ = qc_engine.run_qc(df, columns, client)
    assert "client_brief" not in json.loads(client.beta.messages.requests[0]["messages"][0]["content"])
    assert "QC_Relevant" not in result


def test_extract_setup_cleans_answer():
    client = fake_client()
    brands, topics = qc_engine.extract_setup(client, "the brief text")
    assert client.beta.messages.requests[0]["messages"][0]["content"] == "the brief text"
    assert brands == [
        {"name": "STC", "role": "Client", "aliases": ["اس تي سي", "stc_ksa"]},
        {"name": "Mobily", "role": "Competitor", "aliases": []},
    ]
    assert topics == ["Network", "Pricing"]
